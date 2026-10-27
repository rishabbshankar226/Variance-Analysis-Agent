"""Marcus's holdings scorecard (PLAN.md R18).

A holdings sheet and a consensus panel join on a key column the user names (for example Ticker).
The user also names the one metric that decides a holding (for example Diluted EPS). Each holding's
row for that metric is compared with its consensus by the existing consensus rules, so its label is
that rule's Beat, Miss or In line, and nothing else. Weights are then shared out by label twice: over
the whole book and over the holdings the panel covers, with the uncovered weight shown.
"""

from __future__ import annotations

import csv
import hashlib
import io
from decimal import Decimal

from .intake import CanonicalRole, IntakeResult, _decode_text
from .models import AnalysisResult, ComparisonBasis
from .numeric import financial_context, ratio
from .parsing import parse_decimal


def _fail(message: str) -> IntakeResult:
    return IntakeResult(has_fatal_errors=True, fatal_error_message=message)


def _table(data: bytes, name: str) -> list[dict[str, str]]:
    """A CSV or a workbook's first sheet as rows of text, keyed by the header row."""
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook  # type: ignore[import-untyped]
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        cells = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    else:
        text, _ = _decode_text(data, name)
        cells = list(csv.reader(io.StringIO(text)))
    if not cells:
        return []
    header = [h.strip() for h in cells[0]]
    return [dict(zip(header, (c.strip() for c in row), strict=False)) for row in cells[1:] if any(c.strip() for c in row)]


def _column(rows: list[dict[str, str]], wanted: tuple[str, ...]) -> str | None:
    names = list(rows[0]) if rows else []
    return next((n for n in names if n.casefold() in wanted), None)


def build_scorecard(holdings: bytes, holdings_name: str, panel: bytes, panel_name: str, key: str,
                    metric: str) -> IntakeResult:
    """The covered holdings' deciding-metric rows, ready for the consensus comparison (PLAN.md R18)."""
    try:
        held, offered = _table(holdings, holdings_name), _table(panel, panel_name)
    except ValueError as exc:
        return _fail(str(exc))
    for table, name in ((held, holdings_name), (offered, panel_name)):
        if _column(table, (key.casefold(),)) is None:
            return _fail(f"'{name}' has no '{key}' column to join on.")
    key_h, key_p = _column(held, (key.casefold(),)), _column(offered, (key.casefold(),))
    weight_col = _column(held, ("weight", "actual weight", "weight %"))
    metric_col = _column(offered, ("line item", "metric"))
    if weight_col is None:
        return _fail(f"'{holdings_name}' has no Weight column.")
    if metric_col is None:
        return _fail(f"'{panel_name}' has no Line Item (metric) column.")
    assert key_h is not None and key_p is not None
    marked = {row[weight_col].endswith("%") for row in held if row.get(weight_col)}
    if len(marked) > 1:
        return _fail(f"'{holdings_name}' mixes weights with and without %; give every weight the same way.")
    weights: dict[str, Decimal] = {}
    for row in held:
        holding = row[key_h]
        if holding in weights:
            return _fail(f"'{holding}' appears more than once in '{holdings_name}', so its weight is ambiguous.")
        try:
            weights[holding] = parse_decimal(row[weight_col].rstrip("%"), field="Weight", line_item=holding)
        except ValueError as exc:
            return _fail(str(exc))
    rows: list[dict[str, object]] = []
    seen: set[str] = set()  # the holdings the panel covers
    strangers: list[str] = []
    for row in offered:
        if row.get(metric_col, "").strip() != metric:
            continue
        holding = row[key_p]
        if holding not in weights:
            strangers.append(holding)
            continue
        if holding in seen:
            return _fail(f"'{holding}' has more than one {metric} row in '{panel_name}', so its label is ambiguous.")
        seen.add(holding)
        rows.append({"line_item": holding, "budget": row.get("Consensus", ""), "actual": row.get("Actual", ""),
                     **{k: v for k, v in row.items() if k in {"Measure", "Better", "Consensus Low", "Consensus High",
                                                             "Number of Estimates"} and v}})
    if not rows:
        return _fail(f"No holding in '{holdings_name}' has a {metric} row in '{panel_name}', so nothing can be scored.")
    uncovered = [h for h in weights if h not in seen]
    notes = [f"Joined '{panel_name}' to '{holdings_name}' by '{key}'; each holding is decided by its {metric} row "
             "against consensus"]
    if uncovered:
        notes.append(f"Not covered by the panel for {metric}: " + ", ".join(f"{h} ({_weight(weights[h])})" for h in uncovered))
    if strangers:
        notes.append("Left out of the panel: " + ", ".join(f"{h} (not a holding in '{holdings_name}')" for h in strangers))
    result = IntakeResult(rows=rows, rows_read=len(rows), comparison_basis=ComparisonBasis.CONSENSUS,
                          column_mapping={CanonicalRole.LINE_ITEM: key, CanonicalRole.BUDGET: f"Consensus in {panel_name}",
                                          CanonicalRole.ACTUAL: f"Actual in {panel_name}"},
                          visible_transformations=notes)
    result.scorecard = {"key": key, "metric": metric, "weights": weights, "uncovered": uncovered, "notes": notes,
                        "documents": [{"name": n, "sha256": hashlib.sha256(d).hexdigest(), "role": r}
                                      for n, d, r in ((holdings_name, holdings, "holdings"), (panel_name, panel, "panel"))]}
    return result


def _weight(value: Decimal) -> str:
    from .measures import format_measure
    from .models import MeasureKind
    return format_measure(value, MeasureKind.PERCENT)


def finish_scorecard(intake: IntakeResult, result: AnalysisResult) -> None:
    """Label each holding from the analysis and share the weights out (PLAN.md R18)."""
    from .comparison import display_outcome
    from .report import _rate, _safe_inline
    card = intake.scorecard
    if card is None:
        return
    weights: dict[str, Decimal] = card["weights"]  # type: ignore[assignment]
    labels = {row.line_item: display_outcome(row, ComparisonBasis.CONSENSUS) for row in result.rows}
    book = sum(weights.values(), Decimal(0))
    covered = sum((weights[h] for h in labels), Decimal(0))
    shares: list[tuple[str, Decimal, Decimal | None, Decimal | None]] = []  # label, weight, of book, of covered
    with financial_context():
        for label in ("Beat", "In line", "Miss"):
            weight = sum((weights[h] for h, word in labels.items() if word == label), Decimal(0))
            shares.append((label, weight, ratio(weight, book) if book else None, ratio(weight, covered) if covered else None))
        missing = sum((weights[h] for h in card["uncovered"]), Decimal(0))  # type: ignore[attr-defined]
        shares.append(("Not covered", missing, ratio(missing, book) if book else None, None))
    metric = str(card["metric"])
    intake.report_appendix = [
        f"## Holdings scorecard: {_safe_inline(metric)}", "",
        *(f"- {_safe_inline(str(note))}" for note in card["notes"]), "",  # type: ignore[attr-defined]
        "| Label | Weight | Share of book | Share of covered |", "|---|---:|---:|---:|",
        *(f"| {label} | {_weight(weight)} | {'NM' if of_book is None else _rate(of_book)} | "
          f"{'' if of_covered is None else _rate(of_covered)} |" for label, weight, of_book, of_covered in shares)]
    intake.join = {
        "kind": "scorecard", "key": str(card["key"]), "metric": metric, "documents": card["documents"],
        "holdings": [{"key": h, "weight": format(w, "f"), "label": labels.get(h)} for h, w in weights.items()],
        "shares": [{"label": label, "weight": format(weight, "f"),
                    "share_of_book": None if of_book is None else format(of_book, "f"),
                    "share_of_covered": None if of_covered is None else format(of_covered, "f")}
                   for label, weight, of_book, of_covered in shares],
        "reader_notes": list(card["notes"])}  # type: ignore[call-overload]
