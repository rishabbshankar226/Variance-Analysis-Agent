from __future__ import annotations

import csv
import io
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


_MISSING = {"", "na", "n/a", "none", "null", "-", "—"}


def _require_finite(number: Decimal, *, field: str, line_item: str) -> Decimal:
    if not number.is_finite():
        raise ValueError(f"{line_item}: field '{field}' must be a finite number.")
    return number


def parse_decimal(value: Any, *, field: str, line_item: str) -> Decimal:
    if value is None:
        raise ValueError(f"{line_item}: required field '{field}' is missing.")
    if isinstance(value, Decimal):
        return _require_finite(value, field=field, line_item=line_item)
    if isinstance(value, (int, float)):
        return _require_finite(Decimal(str(value)), field=field, line_item=line_item)

    text = str(value).strip()
    if text.lower() in _MISSING:
        raise ValueError(f"{line_item}: required field '{field}' is missing.")

    negative_parentheses = text.startswith("(") and text.endswith(")")
    if negative_parentheses:
        text = text[1:-1]

    text = (
        text.replace("$", "")
        .replace("€", "")
        .replace("£", "")
        .replace(",", "")
        .replace("_", "")
        .strip()
    )

    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(
            f"{line_item}: field '{field}' is not numeric: {value!r}."
        ) from exc

    if negative_parentheses:
        number = -number
    return _require_finite(number, field=field, line_item=line_item)


def normalize_key(key: str) -> str:
    key = key.strip().lower()
    key = re.sub(r"[^a-z0-9]+", "_", key)
    return key.strip("_")


def _normalized_headers(headers: Iterable[Any]) -> list[str]:
    normalized: list[str] = []
    seen: dict[str, str] = {}
    for raw in headers:
        original = "" if raw is None else str(raw)
        key = normalize_key(original) if raw is not None else ""
        if key and key in seen:
            raise ValueError(
                f"Columns {seen[key]!r} and {original!r} normalize to the same key {key!r}."
            )
        if key:
            seen[key] = original
        normalized.append(key)
    return normalized


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    originals: dict[str, str] = {}
    for raw_key, value in row.items():
        if raw_key is None:
            continue
        original = str(raw_key)
        key = normalize_key(original)
        if key in normalized:
            raise ValueError(
                f"Columns {originals[key]!r} and {original!r} normalize to the same key {key!r}."
            )
        normalized[key] = value
        originals[key] = original
    return normalized


def load_csv_text(text: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV input has no header row.")
    normalized_headers = _normalized_headers(reader.fieldnames)
    output: list[dict[str, Any]] = []
    for row in reader:
        output.append(
            {
                header: row[raw_header]
                for raw_header, header in zip(reader.fieldnames, normalized_headers)
                if header
            }
        )
    return output


def load_json_text(text: str) -> list[dict[str, Any]]:
    payload = json.loads(text)
    if isinstance(payload, dict):
        payload = payload.get("rows", payload.get("data"))
    if not isinstance(payload, list):
        raise ValueError("JSON input must be a list of row objects or contain a 'rows'/'data' list.")
    if not all(isinstance(row, dict) for row in payload):
        raise ValueError("Every JSON row must be an object.")
    return [normalize_row(row) for row in payload]


def load_file(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return load_csv_text(path.read_text(encoding="utf-8-sig"))
    if suffix == ".json":
        return load_json_text(path.read_text(encoding="utf-8"))
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "Excel input requires the optional dependency: pip install 'variance-analysis-agent[excel]'"
            ) from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            raw_headers = list(next(rows))
        except StopIteration as exc:
            raise ValueError("Excel input is empty.") from exc
        headers = _normalized_headers(raw_headers)
        output: list[dict[str, Any]] = []
        for values in rows:
            output.append({h: v for h, v in zip(headers, values) if h})
        return output
    raise ValueError(f"Unsupported input format '{suffix}'. Use CSV, JSON, XLSX, or XLSM.")


def first_present(row: dict[str, Any], candidates: Iterable[str]) -> tuple[str, Any] | None:
    for key in candidates:
        if key in row:
            return key, row[key]
    return None
