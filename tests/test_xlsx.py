"""Formatted Excel workbook export: structure, values, safety and determinism."""

import io
import json
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest

import variance_agent.cli as cli
from variance_agent.analysis import analyze_rows
from variance_agent.models import AnalysisConfig
from variance_agent.xlsx import S_FAVORABLE, S_MONEY, S_PCT, S_UNFAVORABLE, render_xlsx

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "contracts"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHEETS = ("Summary", "Material Variances", "Root Causes & Actions", "All Rows")


def fixture_result(name):
    payload = json.loads((FIXTURES / f"{name}.input.json").read_text(encoding="utf-8"))
    config = payload["config"]
    return analyze_rows(payload["rows"], AnalysisConfig(
        config["period"], Decimal(config["dollar_threshold"]),
        Decimal(config["percent_threshold"]),
    ))


def sheet_cells(data, index):
    """Map cell reference -> (value, style) for one worksheet."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
    cells = {}
    for cell in root.iter(f"{{{NS['m']}}}c"):
        text = cell.find("m:is/m:t", NS)
        number = cell.find("m:v", NS)
        value = text.text if text is not None else number.text if number is not None else None
        cells[cell.get("r")] = (value, int(cell.get("s", "0")))
    return root, cells


@pytest.mark.parametrize("name", ["complete", "incomplete", "zero_budget", "no_material", "summary"])
def test_workbook_parts_are_well_formed_and_deterministic(name):
    result = fixture_result(name)
    data = render_xlsx(result)
    assert data == render_xlsx(result)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        assert names[:5] == ["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                             "xl/_rels/workbook.xml.rels", "xl/styles.xml"]
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())
        for part in names:
            ET.fromstring(archive.read(part))
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    assert tuple(sheet.get("name") for sheet in workbook.iter(f"{{{NS['m']}}}sheet")) == SHEETS


def test_material_sheet_has_numeric_cells_status_colours_and_filters():
    data = render_xlsx(fixture_result("complete"))
    root, cells = sheet_cells(data, 2)
    assert cells["A1"][0] == "Line Item"
    assert cells["A2"] == ("Subscription Revenue", cells["A2"][1])
    assert cells["C2"] == ("1000", S_MONEY)
    assert cells["D2"] == ("1320", S_MONEY)
    assert cells["E2"] == ("320", S_MONEY)
    assert cells["F2"] == ("0.32", S_PCT)
    assert cells["G2"] == ("Favorable", S_FAVORABLE)
    assert cells["G3"] == ("Unfavorable", S_UNFAVORABLE)
    assert root.find("m:sheetViews/m:sheetView/m:pane", NS).get("state") == "frozen"
    assert root.find("m:autoFilter", NS).get("ref") == "A1:I3"


def test_summary_reports_net_impact_only_when_classification_is_complete():
    _, complete = sheet_cells(render_xlsx(fixture_result("complete")), 1)
    net = next(ref for ref, (value, _) in complete.items() if value == "Net Operating Impact")
    row = net[1:]
    assert complete[f"D{row}"][0] == "170"

    _, incomplete = sheet_cells(render_xlsx(fixture_result("incomplete")), 1)
    net = next(ref for ref, (value, _) in incomplete.items() if value == "Net Operating Impact")
    row = net[1:]
    assert [incomplete[f"{col}{row}"][0] for col in "BCDE"] == ["N/A"] * 4
    values = [value for value, _ in incomplete.values()]
    assert "Total Revenue (classified subtotal)" in values
    assert any(value and value.startswith("Net operating impact is N/A") for value in values)


def test_untrusted_labels_are_plain_text_not_markdown_or_markup():
    rows = [
        {"Line Item": "R&D <script>_x_", "Type": "Expense", "Budget": "100", "Actual": "900"},
        {"Line Item": "=HYPERLINK(\"http://x\")", "Type": "Revenue", "Budget": "100",
         "Actual": "900"},
    ]
    result = analyze_rows(rows, AnalysisConfig("Q1", Decimal("10"), Decimal("0.05")))
    data = render_xlsx(result)
    values = [value for value, _ in sheet_cells(data, 2)[1].values()]
    assert "R&D <script>_x_" in values
    assert "=HYPERLINK(\"http://x\")" in values
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        raw = archive.read("xl/worksheets/sheet3.xml").decode("utf-8")
    assert "R&amp;D &lt;script&gt;_x_" in raw
    assert "\\&" not in raw and "**" not in raw
    assert "<f>" not in raw


def test_openpyxl_round_trip_matches_result():
    openpyxl = pytest.importorskip("openpyxl")
    result = fixture_result("complete")
    workbook = openpyxl.load_workbook(io.BytesIO(render_xlsx(result)))
    assert workbook.sheetnames == list(SHEETS)
    sheet = workbook["All Rows"]
    for index, row in enumerate(result.rows, start=2):
        assert sheet.cell(index, 1).value == row.line_item
        assert Decimal(str(sheet.cell(index, 3).value)) == row.budget
        assert Decimal(str(sheet.cell(index, 4).value)) == row.actual
    assert sheet.freeze_panes == "A2"
    assert sheet.cell(2, 3).number_format == "#,##0.00;[Red]-#,##0.00"


def test_cli_writes_workbook(tmp_path):
    source = tmp_path / "input.csv"
    source.write_text("Line Item,Type,Budget,Actual\nSales Revenue,Revenue,100,150\n",
                      encoding="utf-8")
    target = tmp_path / "nested" / "out.xlsx"
    rc = cli.main([str(source), "--period", "Q3 2026", "--dollar-threshold", "10",
                   "--percent-threshold", "5", "--xlsx", str(target)])
    assert rc == 0
    with zipfile.ZipFile(target) as archive:
        assert "xl/worksheets/sheet4.xml" in archive.namelist()


@pytest.mark.parametrize("extra", [["--chat"], ["--check-only"]])
def test_cli_rejects_workbook_with_incompatible_modes(tmp_path, extra, capsys):
    source = tmp_path / "input.csv"
    source.write_text("Line Item,Type,Budget,Actual\nSales Revenue,Revenue,100,150\n",
                      encoding="utf-8")
    args = [str(source), "--period", "Q3", "--dollar-threshold", "10",
            "--percent-threshold", "5", "--xlsx", str(tmp_path / "out.xlsx"), *extra]
    with pytest.raises(SystemExit) as error:
        cli.main(args)
    assert error.value.code == 2
    assert "--xlsx" in capsys.readouterr().err
    assert not (tmp_path / "out.xlsx").exists()


def test_workbook_sentences_spell_out_status_words():
    with zipfile.ZipFile(io.BytesIO(render_xlsx(fixture_result("complete")))) as archive:
        summary = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        actions = archive.read("xl/worksheets/sheet3.xml").decode("utf-8")
    assert "at $320.00 (Favorable)." in summary
    assert "largest material unfavorable variance" in actions
    assert " U variance" not in actions and "(F)" not in summary


def test_markdown_keeps_status_codes():
    from variance_agent.report import render_markdown
    report = render_markdown(fixture_result("complete"))
    assert "at $320.00 (F)." in report
    assert "largest material U variance" in report


# ---------------------------------------------------------------------------
# The redesign changes how the workbook looks, never what it holds
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", ["sample", "consensus", "comparison_set"])
def test_every_value_and_address_is_unchanged_by_the_redesign(name):
    """The snapshots were written by tests/xlsx_values.py from the workbook code at 822e522."""
    pytest.importorskip("openpyxl")
    from tests.xlsx_values import SNAPSHOTS, WORKBOOKS, values
    before = json.loads((SNAPSHOTS / f"{name}.json").read_text(encoding="utf-8"))
    assert values(WORKBOOKS[name]()) == before


def _styles(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("xl/styles.xml").decode("utf-8")


def test_shading_never_carries_meaning():
    styles = _styles(render_xlsx(fixture_result("complete")))
    fills = ET.fromstring(styles).find("m:fills", NS)
    solid = [f.find("m:patternFill/m:fgColor", NS).get("rgb") for f in fills
             if f.find("m:patternFill", NS).get("patternType") == "solid"]
    assert solid == ["FFF2F1EE"]  # the header band only
    xfs = ET.fromstring(styles).find("m:cellXfs", NS)
    for index in (S_FAVORABLE, S_UNFAVORABLE):
        assert xfs[index].get("fillId") == "0"  # the status word carries the meaning, not a fill
    assert "4C1D95" not in styles and "16213E" not in styles  # no violet, no navy


def test_earnings_outcomes_take_their_colour_from_the_word():
    pytest.importorskip("openpyxl")
    import openpyxl
    from tests.xlsx_values import WORKBOOKS
    sheet = openpyxl.load_workbook(io.BytesIO(WORKBOOKS["consensus"]()))["Earnings"]
    heads = [cell.value for cell in sheet[1]]
    col = heads.index("Consensus outcome") + 1
    colours = {sheet.cell(r, col).value: sheet.cell(r, col).font.color.rgb for r in range(2, sheet.max_row + 1)}
    assert colours["Beat"] == "FF175F40" and colours["Miss"] == "FFB03E1B" and colours["In line"] == "FF4A4844"
    actual = heads.index("Actual") + 1
    assert sheet.cell(1, actual).alignment.horizontal == "right" == sheet.cell(2, actual).alignment.horizontal


def test_every_sheet_prints_a_title_page_numbers_and_its_header_row():
    data = render_xlsx(fixture_result("complete"))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
        for index in range(1, 5):
            sheet = archive.read(f"xl/worksheets/sheet{index}.xml").decode("utf-8")
            assert "<oddHeader>&amp;L&amp;\"Calibri,Bold\"Variance report: " in sheet
            assert "Page &amp;P of &amp;N</oddFooter>" in sheet
    assert workbook.count('name="_xlnm.Print_Titles"') == 2  # the two tables with a frozen header
    assert "'Material Variances'!$1:$1" in workbook
