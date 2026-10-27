"""E3: a sheet laid out like EDGAR's Financial_Report.xlsx is read as a financial statement.

Every workbook here is built in the test and every value is fictional.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook

from variance_agent.intake import IntakeResult, read_input_assisted
from variance_agent.models import ComparisonBasis

TITLE = "CONSOLIDATED STATEMENTS OF OPERATIONS - USD ($) shares in Thousands, $ in Millions"
LINES = [("Net sales", 4100, 3900), ("Cost of sales", 2300, 2250), ("Gross margin", 1800, 1650),
         ("Research and development", 300, 280), ("Selling, general and administrative", 250, 240),
         ("Total operating expenses", 550, 520), ("Operating income", 1250, 1130), ("Net income", 1000, 900),
         ("Basic (in dollars per share)", Decimal("6.1"), Decimal("5.5"))]


def _statement(title: str = TITLE, bands: list[tuple[str, int]] | None = None, dates: list[str] | None = None,
               lines: list[tuple[object, ...]] | None = None, sheet_title: str = "Consolidated Statements of Ope") -> bytes:
    """Probe E3's shape: title cell, a merged duration band over date cells, then the lines."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_title
    sheet["A1"] = title
    column = 2
    for band, width in bands if bands is not None else [("12 Months Ended", 2)]:
        sheet.cell(1, column, band)
        if width > 1:
            sheet.merge_cells(start_row=1, start_column=column, end_row=1, end_column=column + width - 1)
        column += width
    for offset, date in enumerate(dates if dates is not None else ["Sep. 27, 2025", "Sep. 28, 2024"]):
        sheet.cell(2, 2 + offset, date)
    for line in lines if lines is not None else LINES:
        sheet.append(list(line))
    capture = io.BytesIO()
    workbook.save(capture)
    return capture.getvalue()


def _read(data: bytes, sheet: str | None = None) -> IntakeResult:
    return read_input_assisted(data, ".xlsx", filename="Financial_Report.xlsx", sheet=sheet)


def test_the_edgar_statement_shape_reads_as_an_income_statement() -> None:
    result = _read(_statement())
    assert not result.has_fatal_errors, result.fatal_error_message
    assert result.comparison_basis is ComparisonBasis.PRIOR_PERIOD
    assert result.comparison_label == "Sep 2024"
    million = Decimal(1_000_000)
    assert [(row["line_item"], row["type"], row["budget"], row["actual"]) for row in result.rows] == [
        ("Net sales", "Revenue", 3900 * million, 4100 * million),
        ("Cost of sales", "Expense", 2250 * million, 2300 * million),
        ("Research and development", "Expense", 280 * million, 300 * million),
        ("Selling, general and administrative", "Expense", 240 * million, 250 * million),
    ]
    notes = result.visible_transformations
    assert notes[0] == ("Read 'Consolidated Statements of Operations' from sheet 'Consolidated Statements of Ope', "
                        "laid out as in an EDGAR Financial_Report.xlsx")
    assert "Read column '12 Months Ended Sep. 27, 2025' as Sep 2025" in notes
    assert ("Left out 'Basic (in dollars per share)' (per-share and share counts are not money totals)") in notes
    assert result.source_kind == "edgar_xlsx"
    assert "Left out 4 subtotals the statement computes from the lines above: Gross margin, Total operating expenses, " \
           "Operating income, Net income" in result.reader_notes


def test_a_balance_sheet_in_this_shape_is_listed_but_not_compared_yet() -> None:
    data = _statement(title="CONSOLIDATED BALANCE SHEETS - USD ($) $ in Millions", bands=[],
                      dates=[], lines=[("", "Sep. 27, 2025", "Sep. 28, 2024"), ("Cash", 100, 90),
                                       ("Total assets", 100, 90)], sheet_title="Consolidated Balance Sheets")
    result = _read(data)
    assert result.has_fatal_errors
    assert result.fatal_error_message == (
        "'Consolidated Balance Sheets' lists assets, liabilities and equity, which can't be compared as "
        "revenue and expense yet.")


def test_a_quarterly_sheet_reads_its_first_duration_and_names_the_other() -> None:
    lines = [("Net sales", 1100, 1000, 3300, 3000), ("Cost of sales", 600, 550, 1800, 1650),
             ("Gross margin", 500, 450, 1500, 1350)]
    data = _statement(bands=[("3 Months Ended", 2), ("9 Months Ended", 2)],
                      dates=["Jun. 28, 2025", "Jun. 29, 2024", "Jun. 28, 2025", "Jun. 29, 2024"], lines=lines)
    result = _read(data)
    assert not result.has_fatal_errors, result.fatal_error_message
    assert [row["actual"] for row in result.rows] == [Decimal(1100_000_000), Decimal(600_000_000)]
    assert ("This sheet also has Consolidated Statements of Operations, 9 Months Ended; choose it as sheet "
            "'Consolidated Statements of Ope (9 Months Ended)'") in result.visible_transformations
    nine = _read(data, sheet="Consolidated Statements of Ope (9 Months Ended)")
    assert not nine.has_fatal_errors, nine.fatal_error_message
    assert [row["actual"] for row in nine.rows] == [Decimal(3300_000_000), Decimal(1800_000_000)]


def test_a_budget_sheet_with_a_statement_title_keeps_todays_read() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in (["Consolidated Statements of Operations - plan review"], ["Line Item", "Budget", "Actual", "Type"],
                ["Revenue", 100, 110, "Revenue"]):
        sheet.append(row)
    capture = io.BytesIO()
    workbook.save(capture)
    result = _read(capture.getvalue())
    assert not result.has_fatal_errors, result.fatal_error_message
    assert result.source_kind is None
    assert result.rows[0]["actual"] == 110


def test_the_renamed_xls_probe_now_reads() -> None:
    result = read_input_assisted(_statement(), ".xls", filename="e8_legacy.xls")
    assert not result.has_fatal_errors, result.fatal_error_message
    assert result.visible_transformations[0] == "Read 'e8_legacy.xls' as an Excel workbook: it is an .xlsx file with an .xls name."


def test_hidden_statement_lines_are_asked_about_counted_whole_or_refused_when_left_out() -> None:
    workbook_bytes = _statement()
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(workbook_bytes))
    workbook.active.row_dimensions[4].hidden = True  # Cost of sales
    capture = io.BytesIO()
    workbook.save(capture)
    data = capture.getvalue()
    asked = read_input_assisted(data, ".xlsx", filename="Financial_Report.xlsx")
    assert asked.needs_hidden_rows_choice
    assert asked.fatal_error_message == ("These rows are hidden in Excel but have amounts, so I need to know whether "
                                         "to count them:\n- Row 4, Cost of sales (hidden)")
    counted = read_input_assisted(data, ".xlsx", filename="Financial_Report.xlsx", hidden_rows="include")
    assert [row["line_item"] for row in counted.rows][:2] == ["Net sales", "Cost of sales"]
    left_out = read_input_assisted(data, ".xlsx", filename="Financial_Report.xlsx", hidden_rows="exclude")
    assert left_out.fatal_error_message == (
        "Sheet 'Consolidated Statements of Ope' has lines hidden in Excel, and leaving lines out would break the "
        "statement's subtotals. Count them, or unhide them, and upload the file again.")
