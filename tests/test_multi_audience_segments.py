"""Flat divisions of fictional Halvorsen Instruments (HALV.FAKE)."""
from dataclasses import replace
from decimal import Decimal
import json
from pathlib import Path

import jsonschema
import pytest

from variance_agent.analysis import analyze_rows
from variance_agent.audience import AudienceConfig
from variance_agent.audit import build_audit_record
from variance_agent.comparisons import analyze_comparisons, render_comparisons
from variance_agent.models import AnalysisConfig, ComparisonBasis
from variance_agent.report import render_markdown


def cfg() -> AudienceConfig:
    return AudienceConfig(period="3Q26", dollar_threshold=Decimal("10"), percent_threshold=Decimal(".02"),
                          comparison_basis=ComparisonBasis.CONSENSUS, extended_headers=True)


def row(name: str, segment: object, amount: str = "100") -> dict[str, object]:
    return {"Line Item": name, "Segment": segment, "Budget": amount, "Actual": amount, "Type": "Revenue"}


def test_segment_duplicates_and_legacy_metadata() -> None:
    rows = [row("Revenue", "North"), row("Revenue", "South")]
    result = analyze_rows(rows, cfg())
    assert [r.segment for r in result.rows] == ["North", "South"]
    assert not any("Duplicate" in w for w in result.warnings)
    assert any("Segment" in w for w in result.warnings)
    assert any("Duplicate" in w for w in analyze_rows([*rows, rows[0]], cfg()).warnings)
    legacy = AnalysisConfig("3Q26", Decimal("10"), Decimal(".02"))
    assert any("Duplicate" in w for w in analyze_rows(rows, legacy).warnings)
    assert all(r.segment == "" for r in analyze_rows(rows, legacy).rows)


@pytest.mark.parametrize("residual,okay", [("-.010001", False), ("-.01", True), ("-.009999", True),
                                             (".009999", True), (".01", True), (".010001", False)])
def test_segment_reconciliation_tolerance(residual: str, okay: bool) -> None:
    rows = [row("Product", "North"), row("Total Revenue", "North", str(Decimal("100") + Decimal(residual))),
            row("Product", "South", "200"), row("Total Revenue", "South", "200")]
    if not okay:
        with pytest.raises(ValueError, match="does not reconcile"):
            analyze_rows(rows, cfg())
        return
    result = analyze_rows(rows, cfg())
    assert [r.summary_verification.detail_count for r in result.rows if r.summary_verification] == [1, 1]
    assert result.revenue.budget == 0  # never synthesize a consolidated total
    record = build_audit_record(result, schema_version="1.8")
    assert [r["revenue"]["budget"] for r in record["segments"]] == ["100", "200"]
    jsonschema.validate(record, json.loads(Path("schemas/audit-1.8.schema.json").read_text()))


def test_scopes_do_not_reconcile_across_segments() -> None:
    rows = [row("Product", "North"), row("Total Revenue", "North", "300"), row("Product", "South", "200")]
    with pytest.raises(ValueError, match="does not reconcile"):
        analyze_rows(rows, cfg())


@pytest.mark.parametrize("segment", [{"North": "Widgets"}, ["North", "Widgets"]])
def test_nested_segments_refused(segment: object) -> None:
    with pytest.raises(ValueError, match="flat segment"):
        analyze_rows([row("Revenue", segment)], cfg())


def test_parent_dimension_refused() -> None:
    with pytest.raises(ValueError, match="flat segment"):
        analyze_rows([{**row("Revenue", "Widgets"), "Parent Segment": "North"}], cfg())


def test_alias_and_safe_report_groups() -> None:
    rows = [{"Line Item": "Revenue", "Region": name, "Consensus": "100", "Actual": "110", "Type": "Revenue"}
            for name in ("North|<x>", "South", "")]
    table = analyze_comparisons(rows, cfg(), (ComparisonBasis.CONSENSUS,))
    assert table.comparisons[0].result.rows[0].segment == "North|<x>"
    report = render_comparisons(table)
    assert "Segment: North\\|&lt;x&gt;" in report
    assert "Segment: South" in report
    assert "Unsegmented" in report
    assert "Region" in " ".join(table.comparisons[0].result.warnings)
    assert render_markdown(table.comparisons[0].result) == report


def test_segment_integrity_verification() -> None:
    from variance_agent.analysis import verify_result
    result = analyze_rows([row("Revenue", "North")], cfg())
    with pytest.raises(RuntimeError, match="segment"):
        verify_result(replace(result, rows=(replace(result.rows[0], segment="South"),)))


def test_comparison_set_segment_join_and_workbook() -> None:
    import io
    from openpyxl import load_workbook
    from variance_agent.xlsx import render_comparison_xlsx
    rows = [{"Line Item": "Revenue", "Segment": segment, "Consensus": value,
             "Our Est.": value, "Actual": "250", "Type": "Revenue"}
            for segment, value in (("North", "100"), ("South", "200"))]
    table = analyze_comparisons(rows, cfg(), (ComparisonBasis.CONSENSUS, ComparisonBasis.ESTIMATE))
    assert [g.segment for g in table.estimate_vs_consensus] == ["North", "South"]
    report = render_comparisons(table)
    assert "North" in report and "South" in report
    wb = load_workbook(io.BytesIO(render_comparison_xlsx(table)))
    assert "Segment" in [c.value for c in wb["Comparisons"][1]]
    single = analyze_comparisons(rows, cfg(), (ComparisonBasis.CONSENSUS,))
    wb = load_workbook(io.BytesIO(render_comparison_xlsx(single)))
    assert "Segment" in [c.value for c in wb["Measures"][1]]
