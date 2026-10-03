import hashlib
import json

import pytest
from openpyxl import Workbook

from variance_agent.cli import main
from variance_agent.analysis import analyze_rows
from variance_agent.audit import build_audit_record
from variance_agent.models import AnalysisConfig
from decimal import Decimal


def args(source):
    return [str(source), "--period", "Q3", "--dollar-threshold", "10", "--percent-threshold", "5"]


def test_check_only_validates_without_report_or_audit_artifacts(tmp_path, capsys):
    source = tmp_path / "source.csv"
    source.write_text("Line Item,Budget,Actual\nSales,100,120\nRent Expense,50,65\n")
    original = source.read_bytes()
    assert main(args(source) + ["--check-only"]) == 0
    output = capsys.readouterr().out
    assert "Validation passed" in output
    assert "2 rows" in output
    assert "# FP&A" not in output
    assert list(tmp_path.iterdir()) == [source]
    assert source.read_bytes() == original


@pytest.mark.parametrize("flag", ["--output", "--audit-json"])
def test_check_only_rejects_output_flags_and_preserves_destination(tmp_path, flag):
    source, destination = tmp_path/"source.csv", tmp_path/"existing.txt"
    source.write_text("Line Item,Budget,Actual\nSales,100,120\n")
    destination.write_text("existing")
    with pytest.raises(SystemExit) as error:
        main(args(source) + ["--check-only", flag, str(destination)])
    assert error.value.code == 2
    assert destination.read_text() == "existing"


def workbook(path):
    book = Workbook()
    first = book.active
    first.title = "Initial"
    first.append(["Line Item", "Budget", "Actual"])
    first.append(["Sales", 100, 110])
    second = book.create_sheet("Reviewed")
    second.append(["Line Item", "Budget", "Actual"])
    second.append(["Sales", 100, 132])
    hidden = book.create_sheet("Hidden inputs")
    hidden.sheet_state = "hidden"
    hidden.append(["Line Item", "Budget", "Actual"])
    hidden.append(["Sales", 100, 140])
    empty = book.create_sheet("Empty")
    empty.append(["Line Item", "Budget", "Actual"])
    book.save(path)
    book.close()


@pytest.mark.parametrize("sheet,actual", [(None,"110"), ("Reviewed","132"), ("Hidden inputs","140")])
def test_selected_worksheet_is_analyzed_and_recorded_with_source_identity(tmp_path, sheet, actual):
    source, audit = tmp_path/"source.xlsx", tmp_path/"audit.json"
    workbook(source)
    options = [] if sheet is None else ["--sheet", sheet]
    assert main(args(source) + options + ["--audit-json", str(audit)]) == 0
    record = json.loads(audit.read_text("utf-8"))
    assert record["totals"]["revenue"]["actual"] == actual
    assert record["source"]["worksheet"] == (sheet or "Initial")
    assert record["source"]["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()


@pytest.mark.parametrize("sheet", ["Missing", "reviewed", "Empty"])
def test_missing_or_empty_sheet_fails_without_replacing_output(tmp_path, sheet):
    source, report = tmp_path/"source.xlsx", tmp_path/"report.md"
    workbook(source)
    report.write_text("existing")
    with pytest.raises(SystemExit) as error:
        main(args(source) + ["--sheet", sheet, "--output", str(report)])
    assert error.value.code == 2
    assert report.read_text() == "existing"


def test_sheet_option_rejects_non_excel_input(tmp_path):
    source = tmp_path/"source.csv"
    source.write_text("Line Item,Budget,Actual\nSales,100,120\n")
    with pytest.raises(SystemExit) as error:
        main(args(source) + ["--sheet", "Reviewed"])
    assert error.value.code == 2


def test_formula_without_cached_result_remains_missing_not_recalculated(tmp_path):
    source = tmp_path/"source.xlsx"
    book = Workbook()
    book.active.append(["Line Item", "Budget", "Actual"])
    book.active.append(["Sales", 100, "=100+20"])
    book.save(source)
    book.close()
    with pytest.raises(SystemExit) as error:
        main(args(source) + ["--check-only"])
    assert error.value.code == 2


def test_identical_financial_facts_on_different_sheets_have_distinct_selection_identity():
    result = analyze_rows([{"Line Item":"Sales", "Budget":"100", "Actual":"132"}],
                          AnalysisConfig("Q3", Decimal("10"), Decimal(".05")))
    first = build_audit_record(result, source_sha256="a"*64, source_worksheet="Initial")
    second = build_audit_record(result, source_sha256="a"*64, source_worksheet="Reviewed")
    assert first["totals"] == second["totals"]
    assert first["source"]["sha256"] == second["source"]["sha256"]
    assert first["provenance"]["configuration_sha256"] != second["provenance"]["configuration_sha256"]
    assert first["analysis_fingerprint"] != second["analysis_fingerprint"]


def test_check_only_preserves_unclassified_warning_behavior(tmp_path, capsys):
    source = tmp_path/"source.csv"
    source.write_text("Line Item,Budget,Actual\nMystery,100,120\n")
    assert main(args(source) + ["--check-only"]) == 0
    output = capsys.readouterr().out
    assert "1 Unclassified" in output
    assert "1 warnings" in output


def test_check_only_rejects_invalid_numeric_input_without_artifacts(tmp_path):
    source = tmp_path/"source.csv"
    source.write_text("Line Item,Budget,Actual\nSales,invalid,120\n")
    with pytest.raises(SystemExit) as error:
        main(args(source) + ["--check-only"])
    assert error.value.code == 2
    assert list(tmp_path.iterdir()) == [source]
