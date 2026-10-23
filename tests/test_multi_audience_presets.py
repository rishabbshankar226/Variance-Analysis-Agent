"""Preset defaults and explicit choices for fictional HALV.FAKE inputs."""
import io
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from variance_agent.cli import main, build_parser

CONSENSUS = "tests/fixtures/multi_audience/consensus.csv"


def flags() -> list[str]:
    return ["--period", "3Q26", "--dollar-threshold", "10", "--percent-threshold", "2", "--per-share-threshold", ".01",
            "--bps-threshold", "25", "--count-threshold", "1000"]


def test_earnings_defaults_and_explicit_overrides(tmp_path: Path) -> None:
    report, workbook, audit = [tmp_path / name for name in ("report.md", "report.xlsx", "audit.json")]
    assert main([CONSENSUS, *flags(), "--preset", "earnings", "-o", str(report), "--xlsx", str(workbook),
                 "--audit-json", str(audit)]) == 0
    text = report.read_text()
    assert text.startswith("# Earnings Review: 3Q26")
    assert text.index("Headline comparisons") < text.index("Other KPIs")
    headline = text.split("## Other KPIs")[0]
    assert "Diluted EPS" in headline and "Gross Margin" in headline and "Revenue" in headline
    assert "Units Shipped" not in headline
    wb = load_workbook(io.BytesIO(workbook.read_bytes()))
    assert "Earnings" in wb.sheetnames and "Measures" in wb.sheetnames
    record = json.loads(audit.read_text())
    assert record["audience"]["preset"] == "earnings"
    assert record["audience"]["layout"] == "earnings"
    assert record["audience"]["per_share_lines"] is True
    assert main([CONSENSUS, *flags(), "--preset", "earnings", "--layout", "standard", "--no-per-share-lines",
                 "-o", str(report), "--audit-json", str(audit)]) == 0
    assert report.read_text().startswith("# Variance Report")
    record = json.loads(audit.read_text())
    assert record["audience"]["layout"] == "standard" and record["audience"]["per_share_lines"] is False


def test_user_selected_headlines_and_no_threshold_defaults(tmp_path: Path) -> None:
    report = tmp_path / "report.md"
    assert main([CONSENSUS, *flags(), "--preset", "earnings", "--headline", "Units Shipped", "-o", str(report)]) == 0
    section = report.read_text().split("## Other KPIs")[0]
    assert "Units Shipped" in section and "Diluted EPS" not in section
    for preset in ("fpa", "earnings", "quant"):
        parsed = build_parser().parse_args(["--preset", preset])
        assert parsed.dollar_threshold is None and parsed.percent_threshold is None
        assert parsed.bps_threshold is None and parsed.per_share_threshold is None and parsed.count_threshold is None


def test_quant_requires_long_output_and_only_explicit_markdown(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as failure:
        main([CONSENSUS, *flags(), "--preset", "quant"])
    assert failure.value.code == 2
    assert "requires --long-output" in capsys.readouterr().err
    long = tmp_path / "long.csv"
    assert main([CONSENSUS, *flags(), "--preset", "quant", "--long-output", str(long)]) == 0
    assert capsys.readouterr().out == ""
    assert len(long.read_text().splitlines()) == 6
    report = tmp_path / "quant.md"
    assert main([CONSENSUS, *flags(), "--preset", "quant", "--long-output", str(long), "-o", str(report)]) == 0
    assert report.exists() and capsys.readouterr().out == ""


def test_explicit_fpa_is_byte_identical_to_no_preset(tmp_path: Path) -> None:
    source = tmp_path / "HALV.FAKE.csv"
    source.write_text("Line Item,Budget,Actual,Type,Region,Consensus\nRevenue,100,110,Revenue,North,999\n")
    outputs = [tmp_path / name for name in ("report.md", "audit.json", "workbook.xlsx")]
    args = [str(source), "--period", "3Q26", "--dollar-threshold", "10", "--percent-threshold", "2",
            "-o", str(outputs[0]), "--audit-json", str(outputs[1]), "--xlsx", str(outputs[2])]
    assert main(args) == 0
    before = [p.read_bytes() for p in outputs]
    assert main([*args, "--preset", "fpa"]) == 0
    assert before == [p.read_bytes() for p in outputs]


def test_prior_year_layout_calls_ratio_yoy(tmp_path: Path) -> None:
    source = tmp_path / "HALV.FAKE.csv"
    source.write_text("Line Item,Type,3Q26,3Q25\nRevenue,Revenue,110,100\n")
    report = tmp_path / "report.md"
    assert main([str(source), "--preset", "earnings", "--compare-to", "prior-period", *flags(), "-o", str(report)]) == 0
    assert "YoY %" in report.read_text()


@pytest.mark.parametrize("extra", [["--long-output"], ["--layout", "earnings", "--long-output"]])
def test_output_or_layout_alone_keeps_legacy_metadata_inert(tmp_path: Path, extra: list[str]) -> None:
    source = tmp_path / "HALV.FAKE.csv"
    source.write_text("Line Item,Budget,Consensus,Model,Actual,Type,Region\nRevenue,100,999,888,110,Revenue,North\n")
    long = tmp_path / "long.csv"
    args = [str(source), "--period", "3Q26", "--dollar-threshold", "10", "--percent-threshold", "2", *extra, str(long)]
    assert main(args) == 0
    import csv
    record = next(csv.DictReader(long.read_text().splitlines()))
    assert record["basis"] == "budget" and record["comparison"] == "100" and record["segment"] == ""


def test_panel_budget_choice_keeps_metadata_columns_inert(tmp_path: Path) -> None:
    import csv
    source = tmp_path / "HALV.FAKE.csv"
    source.write_text("Ticker,Period,Line Item,Budget,Consensus,Model,Actual,Type,Region\n"
                      "HALV.FAKE,3Q26,Revenue,100,999,888,110,Revenue,North\n")
    long, errors = tmp_path / "long.csv", tmp_path / "errors.csv"
    assert main([str(source), "--group-by", "Ticker,Period", "--compare-to", "budget", "--dollar-threshold", "10",
                 "--percent-threshold", "2", "--long-output", str(long), "--errors-output", str(errors)]) == 0
    record = next(csv.DictReader(long.read_text().splitlines()))
    assert record["comparison"] == "100" and record["segment"] == ""
